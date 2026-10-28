import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

def generate_protocol_xlsx(group, members, detailed_results, date_str=None):
    """
    Генерирует протокол комиссии в формате Excel (xlsx).
    group: объект группы (dict)
    members: список сотрудников из группы (для ФИО и профессии)
    detailed_results: список результатов тестов (для оценок)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "СЭД"

    if not date_str:
        date_str = datetime.now().strftime("%d.%m.%Y")

    # --- НАСТРОЙКА СТИЛЕЙ ---
    font_title = Font(name="Times New Roman", size=12, bold=True)
    font_normal = Font(name="Times New Roman", size=11)
    font_bold = Font(name="Times New Roman", size=11, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Ширина колонок
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 35
    ws.column_dimensions['G'].width = 25

    # --- ШАПКА ---
    ws.merge_cells('A1:G1')
    ws['A1'] = f"ПРОТОКОЛ № {group.get('id', '__')}/26 от «{date_str}»"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_center

    ws.merge_cells('A2:G2')
    ws['A2'] = "заседания комиссии по проверке знаний персонала"
    ws['A2'].font = font_normal
    ws['A2'].alignment = align_center

    ws.merge_cells('A3:G3')
    ws['A3'] = f"Группа: {group.get('name', '')} | Обучение: {group.get('description', '').replace('Организация: ', '')}"
    ws['A3'].font = font_bold
    ws['A3'].alignment = align_center
    ws.row_dimensions[3].height = 25

    ws.merge_cells('A4:G4')
    org_name = group.get('organization_name') or "ООО «СЭД»"
    ws['A4'] = f"{org_name}, {group.get('name', '')}"
    ws['A4'].font = font_bold
    ws['A4'].alignment = align_center
    
    ws.merge_cells('A5:G5')
    ws['A5'] = "(Предприятие, подразделение)"
    ws['A5'].font = Font(name="Times New Roman", size=9, italic=True)
    ws['A5'].alignment = align_center

    # --- ТАБЛИЦА (Заголовки) ---
    headers = ["№", "ФИО", "Подразделение", "Профессия, должность", "Общая оценка", "Заключение комиссии", "Роспись об ознакомлении с заключением комиссии"]
    row_offset = 7 
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_offset, column=col_idx, value=header)
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = thin_border
    
    ws.row_dimensions[row_offset].height = 40

    # --- ПОДГОТОВКА ДАННЫХ ---
    # Собираем лучшие результаты по user_id
    results_map = {}
    for r in detailed_results:
        uid = r["user_id"]
        pct = r.get("percentage", 0)
        if uid not in results_map or pct > results_map[uid]:
            results_map[uid] = pct

    # Собираем профессии из текущего состава группы
    members_map = {m["user_id"]: m for m in members}

    # Формируем итоговый список для протокола:
    # Берем всех, кто есть в группе, ПЛЮС тех, кто есть в результатах (но уже удален из группы)
    all_user_ids = set(members_map.keys()) | set(results_map.keys())
    
    # Сортируем по ФИО (для красоты)
    def get_fio(uid):
        if uid in members_map:
            return members_map[uid].get("full_name") or members_map[uid].get("username") or ""
        # Если человека нет в members, ищем его имя в результатах
        for r in detailed_results:
            if r["user_id"] == uid:
                return r.get("full_name") or r.get("username") or "Неизвестный"
        return ""

    sorted_user_ids = sorted(list(all_user_ids), key=get_fio)

    current_row = row_offset + 1
    max_rows = max(24, len(sorted_user_ids)) # Минимум 24 строки
    
    for i in range(max_rows):
        idx = i + 1
        ws.cell(row=current_row, column=1, value=idx).alignment = align_center
        
        if i < len(sorted_user_ids):
            uid = sorted_user_ids[i]
            
            # Данные пользователя
            if uid in members_map:
                m = members_map[uid]
                fio = m.get("full_name") or m.get("username") or ""
                profession = m.get("profession") or "—"
            else:
                fio = get_fio(uid)
                profession = "—" # Если удален из группы, профессии нет
                
            # Оценка
            score = results_map.get(uid)
            if score is not None:
                score_str = f"{float(score):.1f}%"
                conclusion = "допущен к самостоятельной работе" if float(score) >= 60 else "не допущен"
            else:
                score_str = ""
                conclusion = ""
                
            ws.cell(row=current_row, column=2, value=fio).alignment = align_left
            ws.cell(row=current_row, column=3, value=group.get('name', '')).alignment = align_left
            ws.cell(row=current_row, column=4, value=profession).alignment = align_left
            ws.cell(row=current_row, column=5, value=score_str).alignment = align_center
            ws.cell(row=current_row, column=6, value=conclusion).alignment = align_center
        else:
            # Пустые строки для бланка
            for col in range(2, 8):
                ws.cell(row=current_row, column=col, value="")

        for col in range(1, 8):
            ws.cell(row=current_row, column=col).border = thin_border
            ws.cell(row=current_row, column=col).font = font_normal
        
        ws.row_dimensions[current_row].height = 25
        current_row += 1

    # --- ПОДВАЛ (Подписи) ---
    footer_row = current_row + 1
    
    ws.cell(row=footer_row, column=1, value="Председатель комиссии:").font = font_bold
    ws.merge_cells(start_row=footer_row, start_column=2, end_row=footer_row, end_column=4)
    ws.cell(row=footer_row, column=2, value="___________________/Хлыстов В.В./").font = font_normal
    
    ws.cell(row=footer_row, column=5, value="Члены комиссии:").font = font_bold
    
    ws.cell(row=footer_row, column=6, value="1. _________________/Маркина И.А./").font = font_normal
    ws.cell(row=footer_row + 1, column=6, value="2. _________________/Наджафова Н.С./").font = font_normal
    ws.cell(row=footer_row + 2, column=6, value="3. _________________/Якунина Н.С./").font = font_normal

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()